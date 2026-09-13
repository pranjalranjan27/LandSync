# ==============================================================================
# Layer: Core Utilities — Statutory Export Generators (app/core/export_utils.py)
# ALLOWED:
#   - Generate CSV, Excel (.xlsx), and Word (.docx) byte streams from serialized data.
#   - Pure utility functions with zero database models, sessions, or route logic.
# NOT ALLOWED:
#   - Never execute database queries or perform authorization checks here.
# ==============================================================================

import csv
import io
from datetime import date, datetime
from typing import Any, Dict, List, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


CSV_MIME_TYPE = "text/csv; charset=utf-8"
EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def flatten_record(record: Dict[str, Any], parent_key: str = "", sep: str = "_") -> Dict[str, Any]:
    """
    Recursively flattens nested dictionaries and lists into a single-level dictionary
    suitable for tabular presentation in CSV, Excel, and Word.
    """
    items: List[Tuple[str, Any]] = []
    for k, v in record.items():
        new_key = f"{parent_key}{sep}{k}" if parent_key else k
        if isinstance(v, dict):
            items.extend(flatten_record(v, new_key, sep=sep).items())
        elif isinstance(v, list):
            if not v:
                items.append((new_key, ""))
            elif all(isinstance(item, (str, int, float, bool)) for item in v):
                items.append((new_key, ", ".join(str(item) for item in v)))
            elif all(isinstance(item, dict) for item in v):
                # Nested list of dicts: summarize or serialize key identifiers
                summaries = []
                for item in v:
                    summary = item.get("khasra_number") or item.get("name") or item.get("id") or str(item)
                    summaries.append(str(summary))
                items.append((new_key, ", ".join(summaries)))
            else:
                items.append((new_key, ", ".join(str(item) for item in v)))
        elif isinstance(v, (datetime, date)):
            items.append((new_key, v.isoformat()))
        elif isinstance(v, (int, float, bool, str)):
            items.append((new_key, v))
        elif v is None:
            items.append((new_key, ""))
        else:
            items.append((new_key, str(v)))
    return dict(items)


def prepare_tabular_rows(rows: List[Dict[str, Any]]) -> Tuple[List[str], List[Dict[str, Any]]]:
    """
    Normalizes a list of dictionaries by flattening them and compiling a canonical
    header list covering all available keys. Ensures all values are Excel/Word compatible primitives.
    """
    if not rows:
        return [], []

    flattened = [flatten_record(r) for r in rows]
    headers_dict: Dict[str, None] = {}
    for r in flattened:
        for k in r.keys():
            headers_dict[k] = None
    headers = list(headers_dict.keys())

    # Ensure every row has all header keys and values are strictly primitives
    normalized_rows = []
    for r in flattened:
        normalized_row = {}
        for h in headers:
            val = r.get(h, "")
            if isinstance(val, (int, float, bool, str)):
                normalized_row[h] = val
            elif val is None:
                normalized_row[h] = ""
            else:
                normalized_row[h] = str(val)
        normalized_rows.append(normalized_row)

    return headers, normalized_rows


def generate_csv(rows: List[Dict[str, Any]]) -> Tuple[bytes, str]:
    """
    Generates CSV formatted raw bytes with UTF-8 BOM encoding.
    """
    headers, normalized = prepare_tabular_rows(rows)
    output = io.StringIO()

    if headers:
        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for r in normalized:
            writer.writerow(r)
    else:
        output.write("No records available\n")

    csv_bytes = output.getvalue().encode("utf-8-sig")
    return csv_bytes, CSV_MIME_TYPE


def generate_excel(rows: List[Dict[str, Any]], sheet_title: str = "Export") -> Tuple[bytes, str]:
    """
    Generates single-sheet Excel workbook with styled headers and auto-adjusted column widths.
    """
    headers, normalized = prepare_tabular_rows(rows)
    wb = openpyxl.Workbook()
    ws = wb.active

    forbidden_chars = {'[', ']', ':', '*', '?', '/', '\\'}
    clean_title = "".join(c for c in sheet_title if c not in forbidden_chars)[:31] or "Export"
    ws.title = clean_title

    # Header styling (National Government Theme: Dark Slate with White Bold text)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    if headers:
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data rows
        cell_font = Font(name="Calibri", size=10)
        cell_align = Alignment(vertical="center")

        for row_idx, r in enumerate(normalized, start=2):
            row_data = [r.get(h, "") for h in headers]
            ws.append(row_data)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = cell_font
                cell.alignment = cell_align
                cell.border = thin_border

        # Auto-adjust column widths based on content length
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            adjusted_width = max(max_len + 3, 12)
            adjusted_width = min(adjusted_width, 50)  # Bound column width
            ws.column_dimensions[col_letter].width = adjusted_width
    else:
        ws.append(["No records available"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue(), EXCEL_MIME_TYPE


def generate_docx_table(rows: List[Dict[str, Any]], title: str = "Statutory Records Export") -> Tuple[bytes, str]:
    """
    Generates a Word document with a title heading and a clean tabular layout.
    """
    headers, normalized = prepare_tabular_rows(rows)
    doc = docx.Document()

    # Title & Metadata
    title_p = doc.add_heading(title, level=1)
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(12)
    meta_run = meta_p.add_run(
        f"Generated: {datetime.now().strftime('%d %B %Y, %I:%M %p')} | Total Records: {len(normalized)}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    if headers and normalized:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"

        # Format Header Row
        hdr_cells = table.rows[0].cells
        for idx, header in enumerate(headers):
            cell = hdr_cells[idx]
            cell.text = header
            # Header cell shading: Dark Slate #1E293B
            shading_elm = parse_xml(r'<w:shd {} w:fill="1E293B"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading_elm)

            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(255, 255, 255)

        # Populate Data Rows
        for r in normalized:
            row_cells = table.add_row().cells
            for idx, header in enumerate(headers):
                cell = row_cells[idx]
                cell.text = str(r.get(header, ""))
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
    else:
        p = doc.add_paragraph("No records found for the requested criteria.")
        p.runs[0].font.italic = True

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream.getvalue(), DOCX_MIME_TYPE
